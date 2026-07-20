import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from psse_open.spec import SpecWorkbook


class SpecTests(unittest.TestCase):
    def test_hy_spec_reader(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "spec.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Generator Details"
            ws.append(["Parameter", "Value"])
            ws.append(["POC Pbase (MW)", 100.0])
            faults = wb.create_sheet("Fault Types")
            faults.append(["Type", "Code"])
            faults.append(["3PHG", 7])
            study = wb.create_sheet("5255_BalFaults")
            study.append(["File_Name", "PSSE", "Grid_FL_MVA_sig", "Grid_X2R_sig", "Post_Init_Duration_s"])
            study.append(["case_001", True, 300.0, 10.0, 5.0])
            wb.save(str(path))
            spec = SpecWorkbook(str(path))
            scenarios = spec.scenarios()
            self.assertEqual(spec.format, "hy")
            self.assertEqual([scenario.file_name for scenario in scenarios], ["case_001"])
