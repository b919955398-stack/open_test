from __future__ import annotations

import ast
import fnmatch
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence

from openpyxl import load_workbook

from .models import Scenario


class SpecError(ValueError):
    pass


HY_METADATA_SHEETS = {"Generator Details", "Fault Types", "Profiles", "MFRT Sequences"}
LEGACY_METADATA_SHEETS = {"General_information", "Initial Conditions", "grid dist profiles", "MFRT profiles"}


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        if value.lower() in {"true", "yes"}:
            return True
        if value.lower() in {"false", "no"}:
            return False
    return value


def _literal(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    text = value.strip()
    if not text:
        return text
    try:
        return ast.literal_eval(text)
    except (ValueError, SyntaxError):
        return text.strip('"')


def _headers(ws) -> List[str]:
    return [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]


def _records(ws) -> Iterable[tuple]:
    headers = _headers(ws)
    seen: Dict[str, int] = {}
    keys: List[Optional[str]] = []
    for header in headers:
        if not header:
            keys.append(None)
            continue
        count = seen.get(header, 0)
        seen[header] = count + 1
        keys.append(header if count == 0 else "{}__{}".format(header, count + 1))
    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        record = {key: _clean(value) for key, value in zip(keys, row) if key is not None}
        if any(value not in (None, "") for value in record.values()):
            yield row_number, record


def _selected(name: str, patterns: Optional[Sequence[str]]) -> bool:
    return not patterns or any(fnmatch.fnmatchcase(name, pattern) for pattern in patterns)


def _derive_file_name(sheet: str, row_number: int, values: Dict[str, Any]) -> str:
    explicit = values.get("File_Name") or values.get("CSV_name") or values.get("Case_name")
    if explicit not in (None, ""):
        return str(explicit)
    deliverable = str(values.get("Deliverable") or "STUDY")
    clause = values.get("Clause")
    category = str(values.get("Category") or sheet).replace(" ", "")
    test = values.get("Test No") or values.get("TEST_NUMBER") or row_number - 1
    subtest = values.get("Subtest No")
    prefix = "{}_{}_{}".format(deliverable, clause, category) if clause else "{}_{}".format(deliverable, category)
    suffix = "_{:03d}".format(int(test)) if str(test).replace(".", "", 1).isdigit() else "_{}".format(test)
    if subtest not in (None, ""):
        suffix += "p{}".format(subtest)
    return prefix + suffix


class SpecWorkbook:
    """Reads both HY-style clause sheets and the original open DMAT workbook."""

    def __init__(self, path: str):
        self.path = Path(path).resolve()
        if not self.path.exists():
            raise SpecError("SPEC workbook does not exist: {}".format(self.path))
        self._values = load_workbook(str(self.path), data_only=True, read_only=True)
        self._formulas = load_workbook(str(self.path), data_only=False, read_only=True)
        names = set(self._values.sheetnames)
        if "Generator Details" in names:
            self.format = "hy"
        elif "General_information" in names:
            self.format = "legacy"
        else:
            raise SpecError("Unrecognised workbook: expected 'Generator Details' or 'General_information'")
        self.parameters = self._read_parameters()
        self.fault_types = self._read_fault_types()
        self.profiles = self._read_profiles()

    @property
    def sheet_names(self) -> List[str]:
        excluded = HY_METADATA_SHEETS if self.format == "hy" else LEGACY_METADATA_SHEETS
        return [name for name in self._values.sheetnames if name not in excluded]

    def _read_parameters(self) -> Dict[str, Any]:
        sheet = "Generator Details" if self.format == "hy" else "General_information"
        ws = self._values[sheet]
        result: Dict[str, Any] = {}
        if self.format == "hy":
            for _row, record in _records(ws):
                key, value = record.get("Parameter"), record.get("Value")
                if key not in (None, ""):
                    result[str(key)] = _literal(value)
        else:
            for _row, record in _records(ws):
                key, value = record.get("Basic details name"), record.get("Basic details value")
                if key not in (None, ""):
                    result[str(key)] = _literal(value)
        return result

    def _read_fault_types(self) -> Dict[str, int]:
        if self.format == "hy" and "Fault Types" in self._values.sheetnames:
            return {
                str(record["Type"]): int(record["Code"])
                for _row, record in _records(self._values["Fault Types"])
                if record.get("Type") and record.get("Code") is not None
            }
        return {"3PH": 7, "3PHG": 7, "SLG": 1, "1PHG": 1, "2LG": 4, "2PHG": 4, "LL": 8, "1PHPH": 8}

    def _read_profiles(self) -> Dict[str, List[Any]]:
        if self.format != "hy" or "Profiles" not in self._values.sheetnames:
            return {}
        result: Dict[str, List[Any]] = {}
        ws = self._values["Profiles"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] not in (None, ""):
                result[str(row[0]).strip()] = [value for value in row[1:] if value is not None]
        return result

    def scenarios(
        self,
        sheet_patterns: Optional[Sequence[str]] = None,
        enabled_only: bool = True,
        text_filter: Optional[str] = None,
        limit: Optional[int] = None,
    ) -> List[Scenario]:
        scenarios: List[Scenario] = []
        for sheet in self.sheet_names:
            if not _selected(sheet, sheet_patterns):
                continue
            for row_number, values in _records(self._values[sheet]):
                normalized = self._normalize(sheet, values)
                enabled = bool(normalized.get("PSSE", normalized.get("STUDY_EN", False)))
                file_name = _derive_file_name(sheet, row_number, normalized)
                if enabled_only and not enabled:
                    continue
                if text_filter and text_filter.lower() not in (sheet + " " + file_name).lower():
                    continue
                scenarios.append(Scenario(sheet, row_number, file_name, enabled, normalized))
                if limit is not None and len(scenarios) >= limit:
                    return scenarios
        return scenarios

    def _normalize(self, sheet: str, values: Dict[str, Any]) -> Dict[str, Any]:
        result = dict(values)
        if self.format == "legacy":
            aliases = {
                "STUDY_EN": "PSSE",
                "SCR": "Grid_SCR",
                "XR": "Grid_X2R_sig",
                "V": "Vpoc_pu_sig",
                "P": "Ppoc_pu",
                "Q": "Qpoc_pu",
                "Power_base": "Pbase_MW",
                "fstart": "Fault_Time_sig",
                "fdur": "Fault_Duration_sig",
                "zfactor": "Zf2Zs_sig",
                "post_SCR_value": "Post_Fault_SCR",
                "post_XR_value": "Final_Grid_X2R_sig",
                "new_phase": "Phase_Step_deg",
                "rundur": "Post_Init_Duration_s",
                "fault_type": "Fault_Type",
            }
            for source, target in aliases.items():
                if source in result and target not in result:
                    result[target] = result[source]
            pbase = result.get("Power_base", self.parameters.get("Pbase"))
            if result.get("Ppoc_pu") is not None and pbase is not None:
                result["Ppoc_MW_sig"] = float(result["Ppoc_pu"]) * float(pbase)
            if result.get("Qpoc_pu") is not None and pbase is not None:
                result["Qpoc_MVAr_init"] = float(result["Qpoc_pu"]) * float(pbase)
            if result.get("Grid_SCR") is not None:
                base = result.get("SCR_base", self.parameters.get("Mbase", pbase))
                if base is not None:
                    result["Grid_FL_MVA_sig"] = float(result["Grid_SCR"]) * float(base)
            if result.get("Fault_Type") in self.fault_types:
                result["Fault_Type_sig"] = self.fault_types[str(result["Fault_Type"])]
            if sheet == "MFRT" and result.get("MFRT_profile"):
                result["MFRT_profile"] = _literal(result["MFRT_profile"])
            if result.get("profile"):
                result["legacy_profile"] = _literal(result["profile"])
        return result

    def validate_cached_values(self, scenarios: Sequence[Scenario]) -> List[str]:
        warnings: List[str] = []
        if self.format != "hy":
            return warnings
        for scenario in scenarios:
            formula_ws = self._formulas[scenario.sheet]
            value_ws = self._values[scenario.sheet]
            for column in range(1, formula_ws.max_column + 1):
                formula = formula_ws.cell(scenario.row_number, column).value
                value = value_ws.cell(scenario.row_number, column).value
                if isinstance(formula, str) and formula.startswith("=") and value is None:
                    header = formula_ws.cell(1, column).value
                    if header in {"File_Name", "Grid_FL_MVA_sig", "Ppoc_MW_sig", "Qpoc_MVAr_init"}:
                        warnings.append(
                            "{} row {}: formula column {!r} has no cached value; open/save the workbook in Excel".format(
                                scenario.sheet, scenario.row_number, header
                            )
                        )
        return warnings
